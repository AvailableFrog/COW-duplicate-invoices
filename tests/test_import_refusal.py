"""Ticket #3: a workbook that cannot be mapped to the schema is refused, never analysed."""

import subprocess
from pathlib import Path

import openpyxl

from conftest import REPO_ROOT, launch_command, running
from synthetic import SHEET, ticket, without, write_workbook


def launch(workbook: Path, data_dir: Path) -> subprocess.CompletedProcess:
    return subprocess.run(
        launch_command(workbook, "--port", "0", "--data-dir", str(data_dir)),
        capture_output=True, text=True, timeout=20, cwd=REPO_ROOT,
    )


def test_a_workbook_missing_a_required_column_is_refused_naming_the_column(tmp_path):
    workbook = write_workbook(tmp_path / "no-created.xlsx", [ticket(1)], fields=without("Created"))

    result = launch(workbook, tmp_path)

    assert result.returncode == 2
    assert "Listening on" not in result.stdout
    assert "missing required column" in result.stderr
    assert "Created" in result.stderr


def test_an_unreadable_workbook_is_refused_with_a_specific_message(tmp_path):
    workbook = tmp_path / "corrupt.xlsx"
    workbook.write_bytes(b"this is not a zip archive")

    result = launch(workbook, tmp_path)

    assert result.returncode == 2
    assert "could not be read" in result.stderr
    assert "corrupt.xlsx" in result.stderr


def test_a_workbook_without_a_data_sheet_is_refused(tmp_path):
    workbook = tmp_path / "empty.xlsx"
    book = openpyxl.Workbook()
    book.active.title = "Notes"
    book.save(workbook)

    result = launch(workbook, tmp_path)

    assert result.returncode == 2
    assert "no data sheet" in result.stderr
    assert "empty.xlsx" in result.stderr


def test_the_data_sheet_is_the_first_sheet_with_a_header_row(tmp_path):
    workbook = write_workbook(tmp_path / "cover.xlsx", [ticket(1), ticket(2)])
    book = openpyxl.load_workbook(workbook)
    book.create_sheet("Cover", 0)
    book.save(workbook)

    with running(workbook, tmp_path) as booted:
        header = booted.get("/api/snapshot").json()["snapshot"]

    assert header["sheet"] == SHEET
    assert header["row_count"] == 2


def test_a_schema_column_that_appears_twice_is_refused_as_ambiguous(tmp_path):
    workbook = write_workbook(tmp_path / "twice.xlsx", [], fields=[*without(), "Ticket Number"])

    result = launch(workbook, tmp_path)

    assert result.returncode == 2
    assert "more than once" in result.stderr
    assert "Ticket Number" in result.stderr


def test_a_titled_cover_sheet_is_skipped_for_the_sheet_with_the_required_columns(tmp_path):
    workbook = write_workbook(tmp_path / "titled-cover.xlsx", [ticket(1)])
    book = openpyxl.load_workbook(workbook)
    book.create_sheet("Cover", 0).append(["Labour tickets exported for review"])
    book.save(workbook)

    with running(workbook, tmp_path) as booted:
        header = booted.get("/api/snapshot").json()["snapshot"]

    assert header["sheet"] == SHEET
