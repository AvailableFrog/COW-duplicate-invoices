"""Minimal same-schema workbooks built at test time.

The committed export cannot show a missing column, a bad cell or a second
Snapshot, so these tests write a small workbook with the export's 22 columns
into a temporary directory. The release-one spec allows exactly this and
nothing more; the committed synthetic fixture that ADR 0003 defers to release
two (the 50,000-row performance set) is a different thing.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any

import openpyxl

SHEET = "Labour Tickets (1)"

FIELDS = [
    "Employee", "Occupation", "Ticket Number", "Job Name", "Job Number", "Date", "Time",
    "Regular Hours", "Overtime", "Double Time", "Job Work Pay", "Description of Job Work",
    "Supervisor", "Approval Status", "Modified", "Created", "ID", "Email Address", "Keyed",
    "Rate", "Date Approved", "Modified By",
]


class Literal(str):
    """Text written as a plain string cell even when it begins with "=", as typed text in an export is stored.

    openpyxl otherwise writes such a value as a formula; a Literal never is.
    """


def ticket(number: int = 1, **overrides: Any) -> dict[str, Any]:
    """One valid Ticket; pass field values by name with spaces as underscores."""
    values: dict[str, Any] = {
        "Employee": f"Employee {number}",
        "Occupation": "Labourer",
        "Ticket Number": f"L{number:06d}",
        "Job Name": "Beachlands",
        "Job Number": 41215,
        "Date": dt.datetime(2026, 9, 28),
        "Time": "7:30 am to 4:30 pm",
        "Regular Hours": 8,
        "Overtime": 0,
        "Double Time": 0,
        "Job Work Pay": 0,
        "Description of Job Work": "Site work",
        "Supervisor": "Supervisor A",
        "Approval Status": "Approved",
        "Modified": dt.datetime(2026, 10, 5, 7, 48),
        "Created": dt.datetime(2026, 10, 2, 18, 56),
        "ID": 100000 + number,
        "Email Address": f"employee{number}@example.com",
        "Keyed": True,
        "Rate": 14.53,
        "Date Approved": dt.datetime(2026, 10, 5, 7, 42),
        "Modified By": "Reviewer A",
    }
    for key, value in overrides.items():
        name = key.replace("_", " ")
        if name not in values:
            raise KeyError(f"unknown field {name!r}")
        values[name] = value
    return values


def write_workbook(
    path: Path,
    tickets: list[dict[str, Any] | None],
    *,
    fields: list[str] = FIELDS,
    sheet_title: str = SHEET,
) -> Path:
    """Write the header and one row per Ticket; None writes a fully blank row."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = sheet_title
    sheet.append(fields)
    for values in tickets:
        row = [None] * len(fields) if values is None else [values.get(name) for name in fields]
        sheet.append(row)
        for column, value in enumerate(row, start=1):
            if isinstance(value, Literal):
                sheet.cell(row=sheet.max_row, column=column).data_type = "s"
    book.save(path)
    return path


def without(*names: str) -> list[str]:
    return [name for name in FIELDS if name not in names]
